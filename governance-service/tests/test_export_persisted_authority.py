"""Persisted report authority: synthetic run inputs, real DB/API/XLSX bytes."""
from __future__ import annotations

import importlib
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Barrier

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from db import Store
from rule_engine.models import RuleResult, RuleRunResult


def record(label="persisted"):
    return RuleRunResult(
        rule_set="test-rules", version="1", target_summary={"A": 1},
        total=1, passed=0, failed=1, errored=0, score=0, unique_elements=1,
        results=[RuleResult(
            ifc_guid="TEST_GUID", ifc_type="IfcWall", ifc_name="Wall-01",
            rule_code="A", severity="high", status="fail", message=label,
            evidence={"label": label}, usd_prim_path="/World/Wall",
        )],
    )


@pytest.fixture
def module(tmp_path, monkeypatch):
    monkeypatch.setenv("GOV_DB_PATH", str(tmp_path / "governance.db"))
    import app
    return importlib.reload(app)


def seed(module):
    return module.store.create_run("test-v1", "synthetic.ifc", "test-rules")


def export(module, run_id):
    with TestClient(module.app) as client:
        return client.get(f"/api/rule-runs/{run_id}/export?fmt=excel")


def cells(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook}


def test_result_display_fields_survive_store_reopen(module):
    run_id = seed(module)
    module.store.complete_run(run_id, record())
    rows = Store(module.store.db_path).get_results(run_id)
    assert rows[0].get("ifc_type") == "IfcWall"
    assert rows[0].get("ifc_name") == "Wall-01"
    assert rows[0]["ifc_guid"] == "TEST_GUID"


def test_legacy_result_schema_migrates_concurrently_without_inventing_fields(tmp_path):
    path = str(tmp_path / "legacy.db")
    with sqlite3.connect(path) as conn:
        conn.executescript("""
            CREATE TABLE rule_results(
                id TEXT PRIMARY KEY, rule_run_id TEXT, ifc_guid TEXT,
                usd_prim_path TEXT, rule_code TEXT, severity TEXT, status TEXT,
                message TEXT, evidence_json TEXT
            );
            INSERT INTO rule_results VALUES(
                'legacy', 'rr_legacy', 'OLD_GUID', NULL, 'A', 'high', 'fail', 'old', '{}'
            );
        """)
    barrier = Barrier(2)

    def reopen():
        barrier.wait(timeout=5)
        return Store(path).get_results("rr_legacy")

    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(reopen) for _ in range(2)]
        for future in futures:
            row = future.result(timeout=10)[0]
            assert "ifc_type" in row and "ifc_name" in row
            assert row["ifc_type"] is None and row["ifc_name"] is None
            assert row["ifc_guid"] == "OLD_GUID" and row["message"] == "old"
    assert len(Store(path).get_results("rr_legacy")) == 1


def test_export_uses_persisted_winner_not_stale_cache_and_keeps_display_fields(module):
    run_id = seed(module)
    module.store.complete_run(run_id, record())
    module.store.complete_run(run_id, record("late rejected write"))
    module._RUN_CACHE[run_id] = record("stale cache")
    first = cells(export(module, run_id))
    assert first["Failed Elements"][1] == (
        "A", "high", "IfcWall", "Wall-01", "TEST_GUID", "/World/Wall", "persisted"
    )
    module._RUN_CACHE.clear()
    module.store = Store(module.store.db_path)
    assert cells(export(module, run_id)) == first


def test_cache_cannot_export_nonexistent_run(module):
    module._RUN_CACHE["absent"] = record()
    assert export(module, "absent").status_code == 404


@pytest.mark.parametrize("state", ["queued", "running", "failed"])
def test_cache_cannot_export_non_successful_run(module, state):
    run_id = seed(module)
    if state == "running":
        module.store.mark_running(run_id)
    elif state == "failed":
        module.store.fail_run(run_id, "failure")
    module._RUN_CACHE[run_id] = record()
    assert export(module, run_id).status_code == 409


def test_failed_persistence_does_not_publish_computed_report(module, monkeypatch):
    run_id = seed(module)
    monkeypatch.setattr(module, "open_model", lambda _: object())
    monkeypatch.setattr(module, "load_rule_set", lambda _: {})
    monkeypatch.setattr(module, "run_rules", lambda *_: record())
    with sqlite3.connect(module.store.db_path) as conn:
        conn.execute("""
            CREATE TRIGGER reject_results BEFORE INSERT ON rule_results
            BEGIN SELECT RAISE(ABORT, 'test persistence failure'); END
        """)
    module._execute(run_id, "synthetic.ifc", None, None)
    assert module.store.get_run(run_id)["status"] == "failed"
    assert module.store.get_results(run_id) == []
    assert export(module, run_id).status_code == 409
    assert run_id not in module._RUN_CACHE


def test_rebuild_keeps_persisted_identity_and_legacy_missing_identity(module):
    run_id = seed(module)
    run = record()
    run.rule_content_digest = "dsl-json-v1:sha256:" + "a" * 64
    module.store.complete_run(run_id, run)
    module.store = Store(module.store.db_path)
    assert module._rebuild_run_from_store(run_id).rule_content_digest == run.rule_content_digest
    legacy_id = seed(module)
    module.store.complete_run(legacy_id, record())
    with sqlite3.connect(module.store.db_path) as conn:
        conn.execute(
            "UPDATE rule_runs SET summary_json=json_remove(summary_json, '$.rule_content_digest') WHERE id=?",
            (legacy_id,),
        )
    assert module._rebuild_run_from_store(legacy_id).rule_content_digest is None


def test_explicit_pass_evidence_is_readable_after_restart_and_caller_mutation(module):
    run_id = seed(module)
    run = record()
    run.results[0].status = "pass"
    run.passed, run.failed, run.score = 1, 0, 100
    module.store.complete_run(run_id, run)
    run.results[0].evidence["label"] = "mutated"
    run.results[0].status = "fail"
    module.store = Store(module.store.db_path)
    with TestClient(module.app) as client:
        response = client.get(f"/api/rule-runs/{run_id}/results?status=passed")
        assert response.status_code == 200
        rows = response.json()["results"]
        assert len(rows) == 1 and rows[0]["status"] == "pass"
        assert rows[0]["evidence_json"] == '{"label": "persisted"}'
        assert client.get(f"/api/rule-runs/{run_id}/results?status=failed").json()["results"] == []
    assert len(cells(export(module, run_id))["Failed Elements"]) == 1


def test_persisted_spreadsheet_text_never_becomes_a_formula(module):
    run_id = seed(module)
    run = record()
    text = '=HYPERLINK("https://example.invalid","source")'
    run.rule_set = run.version = text
    for field in ("rule_code", "severity", "ifc_type", "ifc_name", "ifc_guid", "usd_prim_path", "message"):
        setattr(run.results[0], field, text)
    module.store.complete_run(run_id, run)
    module.store = Store(module.store.db_path)
    response = export(module, run_id)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert [cell.value for cell in workbook["Failed Elements"][2]] == [text] * 7
    assert workbook["Summary"]["B1"].value == text
    assert workbook["Summary"]["B2"].value == text
    for sheet in workbook:
        for row in sheet:
            assert all(cell.data_type != "f" for cell in row)
